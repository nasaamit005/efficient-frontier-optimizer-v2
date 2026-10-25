
from pathlib import Path
import html
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


# ============================================================
# MARKET & ASSET DIAGNOSTICS
# Creates a separate dashboard and analysis pages.
# Does not modify the Efficient Frontier dashboard.
# ============================================================

ROOT = Path(__file__).resolve().parents[1]
PRICE_FILE = ROOT / "data" / "processed" / "adjusted_prices_clean.csv"
RETURN_FILE = ROOT / "data" / "processed" / "daily_returns.csv"

REPORTS = ROOT / "reports"
PROCESSED = ROOT / "data" / "processed"
DETAILS_DIR = REPORTS / "market_asset_diagnostics_details"

HTML_OUT = REPORTS / "market_asset_diagnostics_dashboard.html"
CSV_OUT = PROCESSED / "market_asset_diagnostics_metrics.csv"
XLSX_OUT = REPORTS / "market_asset_diagnostics_metrics.xlsx"

TRADING_DAYS = 252
RISK_FREE_RATE = 0.02

BG = "#F4F7FB"
PANEL = "#FFFFFF"
INK = "#172B4D"
MUTED = "#61718A"
GRID = "#E5EAF2"
ACCENT = "#4F46E5"

ASSET_COLORS = {
    "AAPL": "#2563EB",
    "MSFT": "#7C3AED",
    "GOOGL": "#0891B2",
    "AMZN": "#EA580C",
    "META": "#DB2777",
    "BND": "#16A34A",
    "^GSPC": "#64748B",
}


# ============================================================
# HELPERS
# ============================================================

def load_frame(path):
    """Load a processed CSV and validate its date column."""
    if not path.exists():
        raise FileNotFoundError(
            f"Required input file was not found:\n{path}"
        )

    df = pd.read_csv(path)

    if "Date" not in df.columns:
        raise ValueError(
            f"'Date' column is missing from {path.name}"
        )

    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df = df.dropna(subset=["Date"])

    if df["Date"].duplicated().any():
        print(
            f"WARNING: duplicate dates in {path.name}; "
            "keeping the first row for each date."
        )

    df = df.drop_duplicates("Date")
    df = df.set_index("Date").sort_index()

    for col in df.columns:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    return df


def pct(value, digits=2):
    if pd.isna(value):
        return "N/A"
    return f"{value:.{digits}%}"


def max_drawdown(prices):
    prices = prices.dropna()
    if prices.empty:
        return np.nan
    return (prices / prices.cummax() - 1).min()


def annualized_cagr(prices):
    prices = prices.dropna()

    if len(prices) < 2:
        return np.nan

    if prices.iloc[0] <= 0 or prices.iloc[-1] <= 0:
        return np.nan

    years = (
        prices.index[-1] - prices.index[0]
    ).days / 365.25

    if years <= 0:
        return np.nan

    return (
        prices.iloc[-1] / prices.iloc[0]
    ) ** (1 / years) - 1


def style(fig, height=450):
    fig.update_layout(
        template="plotly_white",
        paper_bgcolor=PANEL,
        plot_bgcolor=PANEL,
        font=dict(
            family="Arial, sans-serif",
            color=INK,
            size=12,
        ),
        title=dict(
            font=dict(size=18, color=INK),
            x=0.02,
            xanchor="left",
        ),
        height=height,
        margin=dict(l=60, r=35, t=85, b=60),
        hoverlabel=dict(
            bgcolor="#FFFFFF",
            font_size=12,
        ),
        legend=dict(
            orientation="h",
            yanchor="bottom",
            y=1.02,
            xanchor="left",
            x=0,
            bgcolor="rgba(255,255,255,0.85)",
        ),
    )

    fig.update_xaxes(
        showgrid=True,
        gridcolor=GRID,
        zeroline=False,
        linecolor="#D6DEEA",
    )

    fig.update_yaxes(
        showgrid=True,
        gridcolor=GRID,
        zeroline=False,
        linecolor="#D6DEEA",
    )

    return fig


def correlation_summary(block):
    """Mean of the unique off-diagonal correlations."""
    if len(block) < 2:
        return np.nan

    corr = block.corr().to_numpy()
    upper = corr[np.triu_indices_from(corr, k=1)]

    if not len(upper) or np.all(np.isnan(upper)):
        return np.nan

    return float(np.nanmean(upper))


def excel_format_workbook(writer):
    """Apply readable formatting and selected metric color scales."""
    percent_terms = (
        "return",
        "volatility",
        "drawdown",
        "daily_mean",
        "historical_mean",
        "stddev",
    )

    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                "solid",
                fgColor="243B64",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        ws.row_dimensions[1].height = 32

        for column_cells in ws.columns:
            letter = get_column_letter(
                column_cells[0].column
            )

            max_len = max(
                (
                    len(str(cell.value))
                    if cell.value is not None else 0
                )
                for cell in column_cells[:250]
            )

            ws.column_dimensions[letter].width = min(
                max(max_len + 2, 12),
                34,
            )

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")

        headers = {
            str(cell.value): cell.column
            for cell in ws[1]
            if cell.value is not None
        }

        for name, col in headers.items():
            lower_name = name.lower()

            if (
                any(term in lower_name for term in percent_terms)
                or name in {
                    "CAGR",
                    "Total_Return",
                    "Annualized_Mean_Return",
                    "Annualized_Volatility",
                    "Maximum_Drawdown",
                    "Worst_Daily_Return",
                    "Historical_5pct_Return",
                    "Historical_1pct_Return",
                    "Daily_Return",
                }
            ):
                for row_num in range(2, ws.max_row + 1):
                    ws.cell(
                        row_num, col
                    ).number_format = "0.00%"

            if any(
                term in lower_name
                for term in (
                    "sharpe",
                    "sortino",
                    "skewness",
                    "kurtosis",
                    "z_score",
                    "correlation",
                )
            ):
                for row_num in range(2, ws.max_row + 1):
                    ws.cell(
                        row_num, col
                    ).number_format = "0.000"

    ws = writer.book["Asset Metrics"]
    headers = {
        cell.value: cell.column
        for cell in ws[1]
    }

    for name in (
        "Total_Return",
        "CAGR",
        "Annualized_Volatility",
        "Maximum_Drawdown",
    ):
        if name not in headers or ws.max_row < 2:
            continue

        col = headers[name]
        letter = get_column_letter(col)

        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )


# ============================================================
# MAIN
# ============================================================

def main():
    REPORTS.mkdir(parents=True, exist_ok=True)
    PROCESSED.mkdir(parents=True, exist_ok=True)
    DETAILS_DIR.mkdir(parents=True, exist_ok=True)

    print("Loading processed prices and returns...")

    prices_raw = load_frame(PRICE_FILE)
    returns_raw = load_frame(RETURN_FILE)

    assets = [
        asset for asset in prices_raw.columns
        if asset in returns_raw.columns
    ]

    if not assets:
        raise ValueError(
            "No matching asset columns were found "
            "in the price and return files."
        )

    prices_raw = prices_raw[assets]
    returns_raw = returns_raw[assets]

    # Align sources to dates with valid observations for every asset.
    common_dates = prices_raw.index.intersection(
        returns_raw.index
    )

    aligned_prices = prices_raw.loc[common_dates].copy()
    aligned_returns = returns_raw.loc[common_dates].copy()

    valid = (
        aligned_prices.notna().all(axis=1)
        & aligned_returns.notna().all(axis=1)
    )

    prices = aligned_prices.loc[valid].copy()
    returns = aligned_returns.loc[valid].copy()

    if (prices <= 0).any().any():
        raise ValueError(
            "Nonpositive prices were found in the aligned sample. "
            "Review the input data before calculating normalized "
            "performance and CAGR."
        )

    if len(prices) < 252:
        raise ValueError(
            f"Only {len(prices)} aligned observations are available. "
            "At least 252 are recommended."
        )

    n = len(returns)

    print(f"Assets: {', '.join(assets)}")
    print(f"Aligned observations: {n:,}")
    print(
        f"Date range: {returns.index.min().date()} "
        f"to {returns.index.max().date()}"
    )

    # --------------------------------------------------------
    # 1. ASSET METRICS
    # --------------------------------------------------------

    rows = []

    for asset in assets:
        p = prices[asset].dropna()
        r = returns[asset].dropna()

        vol = r.std(ddof=1) * np.sqrt(TRADING_DAYS)
        annual_mean = r.mean() * TRADING_DAYS

        downside = r[r < 0]

        # Downside deviation uses all dates in the denominator.
        downside_deviation = (
            np.sqrt(
                np.mean(np.minimum(r.to_numpy(), 0) ** 2)
                * TRADING_DAYS
            )
            if len(r) else np.nan
        )

        rows.append({
            "Asset": asset,
            "Start_Date": p.index.min().strftime("%Y-%m-%d"),
            "End_Date": p.index.max().strftime("%Y-%m-%d"),
            "Price_Observations": len(p),
            "Return_Observations": len(r),
            "Total_Return": p.iloc[-1] / p.iloc[0] - 1,
            "CAGR": annualized_cagr(p),
            "Annualized_Mean_Return": annual_mean,
            "Annualized_Volatility": vol,
            "Sharpe_Ratio": (
                (annual_mean - RISK_FREE_RATE) / vol
                if pd.notna(vol) and vol > 0 else np.nan
            ),
            "Sortino_Ratio": (
                (annual_mean - RISK_FREE_RATE)
                / downside_deviation
                if pd.notna(downside_deviation)
                and downside_deviation > 0 else np.nan
            ),
            "Worst_Daily_Return": r.min(),
            "Historical_5pct_Return": r.quantile(0.05),
            "Historical_1pct_Return": r.quantile(0.01),
            "Return_Skewness": r.skew(),
            "Return_Excess_Kurtosis": r.kurt(),
            "Maximum_Drawdown": max_drawdown(p),
        })

    metrics = pd.DataFrame(rows)
    metrics.to_csv(CSV_OUT, index=False)

    # --------------------------------------------------------
    # 2. DATA QUALITY AND EXTREME-RETURN FLAGS
    # --------------------------------------------------------

    quality_rows = []
    extreme_rows = []

    for asset in assets:
        r = returns[asset].dropna()
        sd = r.std(ddof=1)

        if pd.notna(sd) and sd > 0:
            z_scores = (r - r.mean()) / sd
            flags = z_scores.abs() > 3
        else:
            z_scores = pd.Series(
                0.0,
                index=r.index,
            )
            flags = pd.Series(
                False,
                index=r.index,
            )

        for date in r.index[flags]:
            extreme_rows.append({
                "Date": date,
                "Asset": asset,
                "Daily_Return": r.loc[date],
                "Historical_Mean": r.mean(),
                "Historical_StdDev": sd,
                "Z_Score": z_scores.loc[date],
                "Flag_Reason": "Absolute z-score exceeds 3",
                "Action": "Review; not automatically removed",
            })

        quality_rows.append({
            "Asset": asset,
            "Aligned_Return_Observations": len(r),
            "Missing_Returns_In_Source": int(
                returns_raw[asset].isna().sum()
            ),
            "Missing_Prices_In_Source": int(
                prices_raw[asset].isna().sum()
            ),
            "Returns_Beyond_3_StdDev": int(flags.sum()),
            "Nonpositive_Prices_In_Source": int(
                (prices_raw[asset] <= 0).sum()
            ),
        })

    quality = pd.DataFrame(quality_rows)
    total_extreme = int(
        quality["Returns_Beyond_3_StdDev"].sum()
    )

    missing_source = int(
        returns_raw[assets].isna().sum().sum()
        + prices_raw[assets].isna().sum().sum()
    )

    extreme_df = pd.DataFrame(extreme_rows)

    if extreme_df.empty:
        extreme_df = pd.DataFrame(columns=[
            "Date",
            "Asset",
            "Daily_Return",
            "Historical_Mean",
            "Historical_StdDev",
            "Z_Score",
            "Flag_Reason",
            "Action",
        ])

    # --------------------------------------------------------
    # 3. HISTORICAL FINDINGS
    # --------------------------------------------------------

    highest = metrics.loc[metrics["Total_Return"].idxmax()]
    lowest = metrics.loc[metrics["Total_Return"].idxmin()]
    deepest = metrics.loc[
        metrics["Maximum_Drawdown"].idxmin()
    ]
    lowest_vol = metrics.loc[
        metrics["Annualized_Volatility"].idxmin()
    ]
    highest_vol = metrics.loc[
        metrics["Annualized_Volatility"].idxmax()
    ]

    corr = returns.corr()

    split1 = max(1, n // 3)
    split2 = max(2, 2 * n // 3)

    blocks = [
        returns.iloc[:split1],
        returns.iloc[split1:split2],
        returns.iloc[split2:],
    ]

    regime_names = [
        "Early period",
        "Middle period",
        "Recent period",
    ]

    regime_corr = [
        correlation_summary(block)
        for block in blocks
    ]

    valid_regimes = [
        i for i, value in enumerate(regime_corr)
        if pd.notna(value)
    ]

    if valid_regimes:
        low_corr_idx = min(
            valid_regimes,
            key=lambda i: regime_corr[i],
        )
        high_corr_idx = max(
            valid_regimes,
            key=lambda i: regime_corr[i],
        )
    else:
        low_corr_idx = 0
        high_corr_idx = 0

    findings = [
        (
            "Performance",
            f"{highest['Asset']} had the highest total return "
            f"({pct(highest['Total_Return'], 1)}); "
            f"{lowest['Asset']} had the lowest "
            f"({pct(lowest['Total_Return'], 1)}). "
            "Historical performance is descriptive, not a forecast."
        ),
        (
            "Risk dispersion",
            f"{lowest_vol['Asset']} had the lowest annualized "
            f"volatility ({pct(lowest_vol['Annualized_Volatility'], 1)}); "
            f"{highest_vol['Asset']} had the highest "
            f"({pct(highest_vol['Annualized_Volatility'], 1)})."
        ),
        (
            "Correlation regimes",
            f"Average pairwise correlation was lowest in the "
            f"{regime_names[low_corr_idx].lower()} "
            f"({regime_corr[low_corr_idx]:.2f}) and highest in the "
            f"{regime_names[high_corr_idx].lower()} "
            f"({regime_corr[high_corr_idx]:.2f}). "
            "These are descriptive chronological periods."
        ),
        (
            "Data review",
            f"The diagnostic flagged {total_extreme:,} asset-day "
            "returns beyond three asset-specific standard deviations. "
            f"Missing entries in inspected source columns: "
            f"{missing_source:,}. No extreme observation was "
            "automatically deleted."
        ),
    ]

    kpis = [
        (
            "Assets analysed",
            str(len(assets)),
            "Equities, bonds and benchmark where available",
        ),
        (
            "Return observations",
            f"{len(returns):,}",
            "Rows after date alignment and validation",
        ),
        (
            "Highest total return",
            highest["Asset"],
            f"{pct(highest['Total_Return'], 1)} over sample",
        ),
        (
            "Deepest drawdown",
            deepest["Asset"],
            f"{pct(deepest['Maximum_Drawdown'], 1)} peak-to-trough",
        ),
    ]

    # --------------------------------------------------------
    # 4. CHART 1: NORMALIZED PERFORMANCE
    # --------------------------------------------------------

    normalized = prices.div(prices.iloc[0]).mul(100)

    fig1 = go.Figure()

    for asset in assets:
        fig1.add_trace(go.Scatter(
            x=normalized.index,
            y=normalized[asset],
            mode="lines",
            name=asset,
            line=dict(
                color=ASSET_COLORS.get(asset, "#2563EB"),
                width=2.2,
            ),
            hovertemplate=(
                f"<b>{asset}</b><br>"
                "Date: %{x|%d %b %Y}<br>"
                "Growth of $100: $%{y:.2f}"
                "<extra></extra>"
            ),
        ))

    fig1.update_layout(
        title="1. Normalized Asset Price Performance",
        xaxis_title="Date",
        yaxis_title="Value of $100 invested",
        hovermode="x unified",
    )
    style(fig1, 470)

    # --------------------------------------------------------
    # 5. CHART 2: ROLLING VOLATILITY
    # --------------------------------------------------------

    rolling = pd.DataFrame(index=returns.index)
    fig2 = go.Figure()

    for asset in assets:
        for window, dash in [(63, "solid"), (126, "dash")]:
            series = (
                returns[asset].rolling(window).std()
                * np.sqrt(TRADING_DAYS)
            )

            rolling[f"{asset}_{window}d"] = series

            fig2.add_trace(go.Scatter(
                x=series.index,
                y=series,
                name=f"{asset} · {window}d",
                mode="lines",
                line=dict(
                    color=ASSET_COLORS.get(asset, "#2563EB"),
                    width=1.8 if window == 63 else 1.3,
                    dash=dash,
                ),
                hovertemplate=(
                    f"<b>{asset}</b> · {window}-day volatility"
                    "<br>Date: %{x|%d %b %Y}"
                    "<br>Annualized volatility: %{y:.2%}"
                    "<extra></extra>"
                ),
            ))

    fig2.update_layout(
        title="2. Rolling Annualized Volatility",
        xaxis_title="Date",
        yaxis_title="Annualized volatility",
        hovermode="x unified",
    )
    fig2.update_yaxes(tickformat=".0%")
    style(fig2, 520)

    # --------------------------------------------------------
    # 6. CHART 3: RETURN DISTRIBUTIONS
    # --------------------------------------------------------

    fig3 = go.Figure()

    for asset in assets:
        fig3.add_trace(go.Histogram(
            x=returns[asset],
            name=asset,
            opacity=0.42,
            nbinsx=90,
            marker_color=ASSET_COLORS.get(asset, "#2563EB"),
            hovertemplate=(
                f"<b>{asset}</b><br>"
                "Daily return: %{x:.2%}"
                "<br>Count: %{y}<extra></extra>"
            ),
        ))

    fig3.update_layout(
        title="3. Daily Return Distributions and Tail Behaviour",
        barmode="overlay",
        xaxis_title="Daily return",
        yaxis_title="Frequency",
    )
    fig3.update_xaxes(tickformat=".0%")
    style(fig3, 470)

    # --------------------------------------------------------
    # 7. CHART 4: CORRELATION REGIMES
    # --------------------------------------------------------

    fig4 = make_subplots(
        rows=1,
        cols=3,
        subplot_titles=regime_names,
        horizontal_spacing=0.08,
    )

    for i, block in enumerate(blocks, start=1):
        c = block.corr()

        fig4.add_trace(go.Heatmap(
            z=c.values,
            x=assets,
            y=assets,
            zmin=-1,
            zmax=1,
            colorscale="RdBu",
            reversescale=True,
            showscale=(i == 3),
            colorbar=dict(
                title="Correlation",
                len=0.8,
            ),
            text=np.round(c.values, 2),
            texttemplate="%{text:.2f}",
            hovertemplate=(
                "Asset X: %{x}<br>"
                "Asset Y: %{y}<br>"
                "Correlation: %{z:.3f}"
                "<extra></extra>"
            ),
        ), row=1, col=i)

    fig4.update_layout(
        title="4. Correlation Regimes Across the Sample",
        height=520,
        paper_bgcolor=PANEL,
        plot_bgcolor=PANEL,
        font=dict(
            family="Arial, sans-serif",
            color=INK,
            size=11,
        ),
        margin=dict(l=65, r=65, t=95, b=65),
    )

    # --------------------------------------------------------
    # 8. CHART 5: DRAWDOWNS
    # --------------------------------------------------------

    drawdowns = prices.div(prices.cummax()).sub(1)
    fig5 = go.Figure()

    for asset in assets:
        fig5.add_trace(go.Scatter(
            x=drawdowns.index,
            y=drawdowns[asset],
            name=asset,
            mode="lines",
            line=dict(
                color=ASSET_COLORS.get(asset, "#2563EB"),
                width=1.7,
            ),
            hovertemplate=(
                f"<b>{asset}</b><br>"
                "Date: %{x|%d %b %Y}<br>"
                "Drawdown: %{y:.2%}<extra></extra>"
            ),
        ))

    fig5.update_layout(
        title="5. Asset-Level Drawdown and Recovery",
        xaxis_title="Date",
        yaxis_title="Drawdown from prior peak",
        hovermode="x unified",
    )
    fig5.update_yaxes(tickformat=".0%")
    style(fig5, 490)

    # --------------------------------------------------------
    # 9. CHART 6: EXTREME RETURNS
    # --------------------------------------------------------

    fig6 = go.Figure()

    for asset in assets:
        r = returns[asset].dropna()
        sd = r.std(ddof=1)

        if pd.notna(sd) and sd > 0:
            flags = (r - r.mean()).abs() > 3 * sd
        else:
            flags = pd.Series(False, index=r.index)

        fig6.add_trace(go.Scatter(
            x=r.index[~flags],
            y=r[~flags],
            mode="markers",
            name=f"{asset} · within 3 SD",
            marker=dict(
                size=3,
                opacity=0.22,
                color=ASSET_COLORS.get(asset, "#2563EB"),
            ),
            showlegend=False,
            hovertemplate=(
                f"<b>{asset}</b><br>"
                "Date: %{x|%d %b %Y}<br>"
                "Daily return: %{y:.2%}<extra></extra>"
            ),
        ))

        if flags.any():
            fig6.add_trace(go.Scatter(
                x=r.index[flags],
                y=r[flags],
                mode="markers",
                name=f"{asset} · extreme",
                marker=dict(
                    size=9,
                    symbol="diamond",
                    color=ASSET_COLORS.get(asset, "#2563EB"),
                    line=dict(color="#FFFFFF", width=1),
                ),
                hovertemplate=(
                    f"<b>{asset} · flagged observation</b>"
                    "<br>Date: %{x|%d %b %Y}"
                    "<br>Daily return: %{y:.2%}<extra></extra>"
                ),
            ))

    fig6.update_layout(
        title="6. Extreme-Return Diagnostics (Beyond 3 Standard Deviations)",
        xaxis_title="Date",
        yaxis_title="Daily return",
        hovermode="closest",
    )
    fig6.update_yaxes(tickformat=".0%")
    style(fig6, 500)

    # --------------------------------------------------------
    # 10. CHART 7: RISK VS RETURN
    # --------------------------------------------------------

    fig7 = go.Figure()

    for _, row in metrics.iterrows():
        asset = row["Asset"]

        fig7.add_trace(go.Scatter(
            x=[row["Annualized_Volatility"]],
            y=[row["Annualized_Mean_Return"]],
            mode="markers+text",
            name=asset,
            text=[asset],
            textposition="top center",
            marker=dict(
                size=15,
                color=ASSET_COLORS.get(asset, "#2563EB"),
                line=dict(color="#FFFFFF", width=1.5),
            ),
            customdata=[[
                row["Sharpe_Ratio"],
                row["Maximum_Drawdown"],
                row["CAGR"],
                row["Total_Return"],
            ]],
            hovertemplate=(
                f"<b>{asset}</b>"
                "<br>Annualized mean return: %{y:.2%}"
                "<br>Annualized volatility: %{x:.2%}"
                "<br>Sharpe ratio: %{customdata[0]:.3f}"
                "<br>Maximum drawdown: %{customdata[1]:.2%}"
                "<br>CAGR: %{customdata[2]:.2%}"
                "<br>Total return: %{customdata[3]:.2%}"
                "<extra></extra>"
            ),
        ))

    fig7.add_hline(
        y=RISK_FREE_RATE,
        line_dash="dash",
        line_color="#64748B",
        annotation_text="2% risk-free-rate reference",
        annotation_position="bottom right",
    )

    fig7.update_layout(
        title="7. Individual Asset Risk vs. Return",
        xaxis_title="Annualized volatility",
        yaxis_title="Annualized arithmetic mean return",
    )
    fig7.update_xaxes(tickformat=".0%")
    fig7.update_yaxes(tickformat=".0%")
    style(fig7, 500)

    # Chart metadata: used by the main dashboard and detail pages.
    chart_specs = [
        (
            "1. Normalized Asset Price Performance",
            fig1,
            "Compare the historical growth of a common starting value.",
        ),
        (
            "2. Rolling Annualized Volatility",
            fig2,
            "Compare changing risk using 63- and 126-observation windows.",
        ),
        (
            "3. Daily Return Distributions and Tail Behaviour",
            fig3,
            "Inspect return dispersion, asymmetry and potential tail risk.",
        ),
        (
            "4. Correlation Regimes Across the Sample",
            fig4,
            "Compare asset co-movement across chronological sample thirds.",
        ),
        (
            "5. Asset-Level Drawdown and Recovery",
            fig5,
            "Compare peak-to-trough losses and historical recovery patterns.",
        ),
        (
            "6. Extreme-Return Diagnostics",
            fig6,
            "Review dates with returns more than three sample standard deviations from the mean.",
        ),
        (
            "7. Individual Asset Risk vs. Return",
            fig7,
            "Compare annualized historical mean return with annualized volatility.",
        ),
    ]

    # --------------------------------------------------------
    # 11. EXCEL WORKBOOK
    # --------------------------------------------------------

    risk_return = metrics[[
        "Asset",
        "Annualized_Mean_Return",
        "Annualized_Volatility",
        "Sharpe_Ratio",
        "Sortino_Ratio",
        "CAGR",
        "Total_Return",
        "Maximum_Drawdown",
    ]].copy()

    methodology = pd.DataFrame([
        [
            "Input data",
            "Existing processed adjusted prices and daily returns",
        ],
        [
            "Date alignment",
            "Common dates with valid observations across included assets",
        ],
        [
            "Annualization",
            "Mean daily return × 252; daily volatility × sqrt(252)",
        ],
        [
            "CAGR",
            "Price-based compounded annual growth over observed calendar time",
        ],
        [
            "Sharpe ratio",
            "(Annualized arithmetic mean return − assumed 2% annual risk-free rate) / annualized volatility",
        ],
        [
            "Sortino ratio",
            "Excess annualized arithmetic mean divided by annualized downside deviation",
        ],
        [
            "Downside deviation",
            "Square root of 252 times the average squared negative daily return, including zero contributions from nonnegative returns",
        ],
        [
            "Rolling volatility",
            "Rolling sample standard deviation × sqrt(252); 63- and 126-observation windows",
        ],
        [
            "Normalized performance",
            "Price / first aligned price × 100",
        ],
        [
            "Drawdown",
            "Price / running maximum price − 1",
        ],
        [
            "Correlation",
            "Pearson correlation of daily returns",
        ],
        [
            "Correlation periods",
            "Three chronological blocks, approximately equal in observation count",
        ],
        [
            "Extreme-return flag",
            "Absolute deviation from asset sample mean greater than 3 sample standard deviations",
        ],
        [
            "Extreme handling",
            "Diagnostic flags only; no extreme observations automatically deleted",
        ],
        [
            "Risk-free rate",
            "Assumed annual rate of 2%; not a live market quote",
        ],
        [
            "Trading-day convention",
            "252 observations per year",
        ],
        [
            "Limitations",
            "Historical estimates depend on the sample and do not guarantee future outcomes",
        ],
    ], columns=["Topic", "Definition_or_Assumption"])

    with pd.ExcelWriter(
        XLSX_OUT,
        engine="openpyxl",
    ) as writer:
        metrics.to_excel(
            writer,
            sheet_name="Asset Metrics",
            index=False,
        )

        quality.to_excel(
            writer,
            sheet_name="Data Quality",
            index=False,
        )

        prices.to_excel(
            writer,
            sheet_name="Aligned Prices",
        )

        returns.to_excel(
            writer,
            sheet_name="Daily Returns",
        )

        normalized.to_excel(
            writer,
            sheet_name="Normalized Performance",
        )

        rolling.to_excel(
            writer,
            sheet_name="Rolling Volatility",
        )

        drawdowns.to_excel(
            writer,
            sheet_name="Drawdowns",
        )

        corr.to_excel(
            writer,
            sheet_name="Full Correlation",
        )

        risk_return.to_excel(
            writer,
            sheet_name="Risk Return",
            index=False,
        )

        extreme_df.to_excel(
            writer,
            sheet_name="Extreme Observations",
            index=False,
        )

        methodology.to_excel(
            writer,
            sheet_name="Methodology",
            index=False,
        )

        for sheet_name, block in zip(
            [
                "Early Correlation",
                "Middle Correlation",
                "Recent Correlation",
            ],
            blocks,
        ):
            block.corr().to_excel(
                writer,
                sheet_name=sheet_name,
            )

        pd.DataFrame({
            "Period": regime_names,
            "Average_Pairwise_Correlation": regime_corr,
            "Observations": [
                len(block) for block in blocks
            ],
            "Start_Date": [
                block.index.min().strftime("%Y-%m-%d")
                if not block.empty else "N/A"
                for block in blocks
            ],
            "End_Date": [
                block.index.max().strftime("%Y-%m-%d")
                if not block.empty else "N/A"
                for block in blocks
            ],
        }).to_excel(
            writer,
            sheet_name="Correlation Regimes",
            index=False,
        )

        excel_format_workbook(writer)

    # --------------------------------------------------------
    # 12. INTERACTIVE METRICS TABLE
    # --------------------------------------------------------

    table_headers = list(metrics.columns)

    header_html = "".join(
        f'<th data-col="{i}" tabindex="0" title="Click to sort">'
        f'{html.escape(str(col))} '
        '<span class="sort-mark">↕</span></th>'
        for i, col in enumerate(table_headers)
    )

    table_rows = []

    percentage_columns = {
        "Total_Return",
        "CAGR",
        "Annualized_Mean_Return",
        "Annualized_Volatility",
        "Worst_Daily_Return",
        "Historical_5pct_Return",
        "Historical_1pct_Return",
        "Maximum_Drawdown",
    }

    ratio_columns = {
        "Sharpe_Ratio",
        "Sortino_Ratio",
        "Return_Skewness",
        "Return_Excess_Kurtosis",
    }

    for _, row in metrics.iterrows():
        cells = []

        for col in table_headers:
            value = row[col]

            if pd.isna(value):
                display = "N/A"
            elif col in percentage_columns:
                display = f"{value:.2%}"
            elif col in ratio_columns:
                display = f"{value:.3f}"
            else:
                display = str(value)

            cells.append(
                f'<td data-value="{html.escape(str(value))}">'
                f'{html.escape(display)}</td>'
            )

        table_rows.append(
            f'<tr data-asset="{html.escape(str(row["Asset"]))}">'
            + "".join(cells)
            + "</tr>"
        )

    table_html = (
        '<table id="metricsTable">'
        '<thead><tr>'
        + header_html
        + '</tr></thead><tbody>'
        + "".join(table_rows)
        + '</tbody></table>'
    )

    quality_html = quality.to_html(
        index=False,
        classes="quality-table",
        border=0,
    )

    # --------------------------------------------------------
    # 13. DEDICATED CHART ANALYSIS PAGES
    # --------------------------------------------------------

    chart_details = [
        {
            "slug": "01_normalized_performance",
            "what": "Compares the historical growth of a common starting value of 100 for every asset.",
            "why": "Makes price growth comparable across assets with different starting prices.",
            "formula": "Normalized value = price at date / first aligned price × 100.",
            "finding": (
                f"{highest['Asset']} had the highest total return "
                f"({pct(highest['Total_Return'])}); "
                f"{lowest['Asset']} had the lowest "
                f"({pct(lowest['Total_Return'])})."
            ),
            "limitations": "Price growth is not a forecast. Interpretation as investor total return depends on whether the input adjusted prices correctly reflect distributions and corporate actions.",
        },
        {
            "slug": "02_rolling_volatility",
            "what": "Shows annualized rolling volatility for 63- and 126-observation windows.",
            "why": "Reveals how historical return variability changes over time.",
            "formula": "Rolling volatility = rolling sample standard deviation of daily returns × sqrt(252).",
            "finding": (
                f"{lowest_vol['Asset']} had the lowest full-sample "
                f"annualized volatility "
                f"({pct(lowest_vol['Annualized_Volatility'])}); "
                f"{highest_vol['Asset']} had the highest "
                f"({pct(highest_vol['Annualized_Volatility'])})."
            ),
            "limitations": "Results depend on window length and the sample. Volatility does not capture every form of investment risk.",
        },
        {
            "slug": "03_return_distribution",
            "what": "Compares the frequency and shape of daily asset returns.",
            "why": "Helps reveal dispersion and tail behavior that averages can conceal.",
            "formula": "Histograms group daily returns into bins; skewness, excess kurtosis and empirical quantiles are calculated from observed returns.",
            "finding": "Use the Asset Metrics sheet to compare skewness, excess kurtosis, worst daily return and historical 1% and 5% quantiles for each asset.",
            "limitations": "Histogram shape depends on binning. Tail quantiles and higher moments are uncertain estimates, especially with limited tail observations.",
        },
        {
            "slug": "04_correlation_regimes",
            "what": "Compares Pearson correlations across three chronological periods.",
            "why": "Shows how historical co-movement may vary across the sample and informs diversification analysis.",
            "formula": "Pearson correlation = covariance of two return series / product of their standard deviations.",
            "finding": (
                f"Average pairwise correlation was lowest in the "
                f"{regime_names[low_corr_idx].lower()} "
                f"({regime_corr[low_corr_idx]:.3f}) and highest in the "
                f"{regime_names[high_corr_idx].lower()} "
                f"({regime_corr[high_corr_idx]:.3f})."
            ),
            "limitations": "These are chronological thirds, not statistically estimated market regimes. Correlation does not establish causation.",
        },
        {
            "slug": "05_drawdown_recovery",
            "what": "Tracks each asset's decline from its own previous running price peak.",
            "why": "Shows the depth and timing of historical losses beyond volatility alone.",
            "formula": "Drawdown at date t = price at t / highest price observed through t − 1.",
            "finding": (
                f"{deepest['Asset']} had the deepest observed maximum "
                f"drawdown ({pct(deepest['Maximum_Drawdown'])})."
            ),
            "limitations": "Drawdowns depend on sample start and end dates. Historical drawdown does not quantify the probability of future losses.",
        },
        {
            "slug": "06_extreme_returns",
            "what": "Flags returns whose absolute distance from their asset's sample mean exceeds three sample standard deviations.",
            "why": "Identifies dates for data-quality review and potential event investigation.",
            "formula": "Z-score = (daily return − sample mean) / sample standard deviation; flag when absolute Z-score exceeds 3.",
            "finding": (
                f"The rule flagged {total_extreme:,} asset-day observations. "
                "Flags are not proof of errors."
            ),
            "limitations": "Financial returns may have heavy tails and changing volatility. This screen is not a universal outlier test. Observations are not automatically deleted.",
        },
        {
            "slug": "07_risk_return",
            "what": "Plots annualized arithmetic mean return against annualized volatility for each asset.",
            "why": "Provides a descriptive comparison of historical return and variability.",
            "formula": "Annualized mean = mean daily return × 252; annualized volatility = sample standard deviation × sqrt(252).",
            "finding": (
                f"{metrics.loc[metrics['Sharpe_Ratio'].idxmax(), 'Asset']} "
                "had the highest sample Sharpe ratio using the assumed "
                "2% annual risk-free rate."
            ),
            "limitations": "Historical means are noisy. Single-asset comparisons are not optimized portfolio results and do not imply an investment recommendation.",
        },
    ]

    # Write each detail page.
    for (title, fig, note), detail in zip(
        chart_specs,
        chart_details,
    ):
        detail_chart_html = fig.to_html(
            full_html=False,
            include_plotlyjs="cdn",
            config={
                "responsive": True,
                "displaylogo": False,
                "scrollZoom": True,
                "toImageButtonOptions": {
                    "format": "png",
                    "scale": 4,
                    "filename": detail["slug"],
                },
            },
        )

        analysis_table = metrics.to_html(
            index=False,
            classes="analysis-table",
            border=0,
            float_format=lambda x: f"{x:.4f}",
        )

        detail_page = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)} — Full Analysis</title>
<style>
* {{ box-sizing: border-box; }}
body {{
  margin: 0;
  background: #F4F7FB;
  color: #172B4D;
  font-family: Inter, Segoe UI, Arial, sans-serif;
  line-height: 1.65;
}}
header {{
  padding: 28px clamp(18px, 4vw, 55px);
  color: white;
  background: linear-gradient(115deg, #173B70, #4F46E5 58%, #0EA5A4);
}}
header h1 {{
  margin: 8px 0;
  font-size: clamp(25px, 3vw, 36px);
}}
header p {{ max-width: 950px; margin: 0; }}
main {{
  max-width: 1400px;
  margin: auto;
  padding: 24px clamp(12px, 3vw, 38px) 50px;
}}
.panel {{
  background: white;
  border: 1px solid #E4EAF3;
  border-radius: 16px;
  padding: 20px;
  margin: 18px 0;
  box-shadow: 0 5px 18px #1E3A5F08;
}}
h2 {{ margin: 0 0 8px; font-size: 21px; }}
.muted {{ color: #61718A; }}
.finding {{
  border-left: 5px solid #0F766E;
  background: #F0FDFA;
  padding: 14px 17px;
  border-radius: 8px;
}}
a {{ color: #4338CA; font-weight: 700; }}
.analysis-table {{
  border-collapse: collapse;
  width: 100%;
  font-size: 12px;
}}
.analysis-table th {{
  background: #243B64;
  color: white;
  text-align: left;
}}
.analysis-table th, .analysis-table td {{
  padding: 9px;
  border-bottom: 1px solid #E5EAF2;
}}
.table-wrap {{ overflow: auto; }}
@media(max-width: 650px) {{
  .panel {{ padding: 12px; }}
}}
</style>
</head>
<body>
<header>
  <p>MARKET &amp; ASSET DIAGNOSTICS · FULL ANALYSIS</p>
  <h1>{html.escape(title)}</h1>
  <p>{html.escape(note)}</p>
</header>
<main>
  <p><a href="../market_asset_diagnostics_dashboard.html">← Return to main dashboard</a></p>

  <section class="panel">
    <h2>Interactive chart</h2>
    <p class="muted">Hover for exact values; use zoom, pan, legend controls, and the camera icon for a high-resolution export.</p>
    {detail_chart_html}
  </section>

  <section class="panel">
    <h2>What this chart measures</h2>
    <p>{html.escape(detail["what"])}</p>
    <h2>Why it matters for portfolio research</h2>
    <p>{html.escape(detail["why"])}</p>
    <h2>Construction and formula</h2>
    <p>{html.escape(detail["formula"])}</p>
  </section>

  <section class="panel">
    <h2>Calculated finding</h2>
    <div class="finding">{html.escape(detail["finding"])}</div>
  </section>

  <section class="panel">
    <h2>Assumptions and limitations</h2>
    <p>{html.escape(detail["limitations"])}</p>
    <p class="muted">Annualization convention: {TRADING_DAYS} observations per year. Assumed annual risk-free rate: {RISK_FREE_RATE:.2%}. Results describe the supplied historical sample and are not forecasts or investment advice.</p>
  </section>

  <section class="panel">
    <h2>Asset-level numerical results</h2>
    <p class="muted">The table provides calculated metrics for all included assets. The Excel workbook contains the daily series and chart-specific data.</p>
    <div class="table-wrap">{analysis_table}</div>
  </section>
</main>
</body>
</html>
"""

        (DETAILS_DIR / f"{detail['slug']}.html").write_text(
            detail_page,
            encoding="utf-8",
        )

    # --------------------------------------------------------
    # 14. MAIN DASHBOARD CHART CARDS
    # --------------------------------------------------------

    chart_sections = []

    for i, (title, fig, note) in enumerate(chart_specs):
        detail = chart_details[i]

        detail_link = (
            "market_asset_diagnostics_details/"
            f"{detail['slug']}.html"
        )

        chart_html = fig.to_html(
            full_html=False,
            include_plotlyjs=(i == 0),
            config={
                "responsive": True,
                "displaylogo": False,
                "scrollZoom": True,
                "toImageButtonOptions": {
                    "format": "png",
                    "scale": 3,
                    "filename": f"market_diagnostics_chart_{i + 1}",
                },
            },
        )

        chart_sections.append(f"""
        <section class="chart-card">
          <div class="chart-note">
            <b>{html.escape(title)}</b>
            <p>{html.escape(note)}</p>
            <a class="analysis-link" href="{detail_link}">
              View Full Analysis →
            </a>
          </div>
          {chart_html}
        </section>
        """)

    kpi_html = ""

    for i, (label, value, sub) in enumerate(kpis):
        kpi_html += f"""
        <div class="kpi kpi-{i+1}">
          <div class="kpi-label">{html.escape(label)}</div>
          <div class="kpi-value">{html.escape(value)}</div>
          <div class="kpi-sub">{html.escape(sub)}</div>
        </div>
        """

    finding_html = ""

    for i, (title, body) in enumerate(findings):
        finding_html += f"""
        <article class="finding finding-{i+1}">
          <div class="finding-index">0{i+1}</div>
          <div>
            <h3>{html.escape(title)}</h3>
            <p>{html.escape(body)}</p>
          </div>
        </article>
        """

    # --------------------------------------------------------
    # 15. FINAL HTML PAGE
    # --------------------------------------------------------

    page = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Market &amp; Asset Diagnostics</title>
<style>
:root {{
  --bg: {BG};
  --panel: {PANEL};
  --ink: {INK};
  --muted: {MUTED};
  --grid: {GRID};
  --accent: {ACCENT};
}}
* {{ box-sizing: border-box; }}
body {{
  margin: 0;
  background: var(--bg);
  color: var(--ink);
  font-family: Inter, Segoe UI, Arial, sans-serif;
}}
header {{
  padding: 32px clamp(18px, 4vw, 56px);
  color: white;
  background: linear-gradient(115deg, #173B70, #4F46E5 58%, #0EA5A4);
}}
header .eyebrow {{
  text-transform: uppercase;
  letter-spacing: 2px;
  font-size: 12px;
  opacity: .85;
  font-weight: 700;
}}
header h1 {{
  margin: 10px 0 8px;
  font-size: clamp(27px, 4vw, 42px);
}}
header p {{
  margin: 0;
  max-width: 950px;
  line-height: 1.6;
  opacity: .93;
}}
main {{
  max-width: 1600px;
  margin: auto;
  padding: 24px clamp(12px, 3vw, 40px) 55px;
}}
.section-title {{
  display: flex;
  align-items: end;
  justify-content: space-between;
  gap: 12px;
  flex-wrap: wrap;
  margin: 26px 0 15px;
}}
.section-title h2 {{ margin: 0; font-size: 23px; }}
.section-title p {{ color: var(--muted); margin: 6px 0 0; }}
.kpi-grid {{
  display: grid;
  grid-template-columns: repeat(4, minmax(0, 1fr));
  gap: 15px;
}}
.kpi {{
  position: relative;
  overflow: hidden;
  background: white;
  border: 1px solid #E4EAF3;
  border-radius: 16px;
  padding: 20px;
  box-shadow: 0 5px 18px #1E3A5F0A;
}}
.kpi:before {{
  content: "";
  position: absolute;
  left: 0;
  top: 0;
  bottom: 0;
  width: 5px;
  background: var(--kpi);
}}
.kpi-1 {{ --kpi: #2563EB; }}
.kpi-2 {{ --kpi: #0891B2; }}
.kpi-3 {{ --kpi: #16A34A; }}
.kpi-4 {{ --kpi: #DB2777; }}
.kpi-label {{ color: var(--muted); font-weight: 650; font-size: 13px; }}
.kpi-value {{
  font-size: clamp(24px, 2.5vw, 34px);
  font-weight: 800;
  margin: 11px 0 5px;
  overflow-wrap: anywhere;
}}
.kpi-sub {{ font-size: 12px; color: var(--muted); line-height: 1.45; }}
.findings-grid {{
  display: grid;
  grid-template-columns: repeat(2, minmax(0, 1fr));
  gap: 14px;
}}
.finding {{
  display: flex;
  gap: 15px;
  background: white;
  padding: 19px;
  border-radius: 15px;
  border: 1px solid #E4EAF3;
}}
.finding-index {{
  display: grid;
  place-items: center;
  flex: 0 0 39px;
  height: 39px;
  border-radius: 12px;
  color: white;
  font-weight: 800;
  background: var(--fc);
}}
.finding-1 {{ --fc: #2563EB; }}
.finding-2 {{ --fc: #D97706; }}
.finding-3 {{ --fc: #0891B2; }}
.finding-4 {{ --fc: #7C3AED; }}
.finding h3 {{ margin: 0 0 7px; font-size: 16px; }}
.finding p {{
  margin: 0;
  color: var(--muted);
  line-height: 1.6;
  font-size: 13px;
}}
.chart-grid {{
  display: grid;
  grid-template-columns: repeat(2, minmax(0, 1fr));
  gap: 16px;
}}
.chart-card {{
  background: white;
  border: 1px solid #E4EAF3;
  border-radius: 16px;
  overflow: hidden;
  box-shadow: 0 5px 18px #1E3A5F08;
  min-width: 0;
}}
.chart-card:last-child {{ grid-column: 1 / -1; }}
.chart-note {{ padding: 18px 20px 0; font-size: 13px; line-height: 1.55; }}
.chart-note p {{ margin: 5px 0 0; color: var(--muted); }}

.analysis-link {{
  display: inline-block;
  margin: 12px 0 4px;
  padding: 9px 13px;
  border-radius: 8px;
  background: #EEF2FF;
  color: #4338CA;
  font-weight: 750;
  text-decoration: none;
}}
.analysis-link:hover {{
  background: #DDE6FF;
}}

.table-panel {{
  background: white;
  border: 1px solid #E4EAF3;
  border-radius: 16px;
  padding: 18px;
  overflow: hidden;
}}
.table-controls {{
  display: flex;
  gap: 10px;
  flex-wrap: wrap;
  align-items: center;
  margin: 0 0 15px;
}}
.table-controls input, .table-controls select {{
  border: 1px solid #CBD5E1;
  background: #fff;
  border-radius: 9px;
  padding: 11px 12px;
  color: var(--ink);
}}
.table-controls input {{ flex: 1; min-width: 210px; }}
.action-btn {{
  border: 0;
  border-radius: 9px;
  padding: 11px 15px;
  cursor: pointer;
  color: white;
  font-weight: 700;
  background: #4F46E5;
}}
.action-btn.secondary {{ background: #0F766E; }}
.table-scroll {{
  overflow: auto;
  max-height: 620px;
  border: 1px solid #E5EAF2;
  border-radius: 10px;
}}
#metricsTable {{
  border-collapse: separate;
  border-spacing: 0;
  width: 100%;
  white-space: nowrap;
  font-size: 12px;
}}
#metricsTable th {{
  position: sticky;
  top: 0;
  z-index: 2;
  background: #243B64;
  color: white;
  text-align: left;
  padding: 12px 10px;
  cursor: pointer;
}}
#metricsTable th:hover {{ background: #344F82; }}
#metricsTable td {{
  padding: 10px;
  border-bottom: 1px solid #EDF1F7;
}}
#metricsTable tbody tr:nth-child(even) {{ background: #F7F9FC; }}
#metricsTable tbody tr:hover {{ background: #E9F2FF; }}
#tableCount {{ color: var(--muted); font-size: 12px; margin: 12px 0 0; }}
.quality-wrap {{
  overflow: auto;
  background: white;
  border-radius: 14px;
  padding: 15px;
  border: 1px solid #E4EAF3;
}}
.quality-table {{
  width: 100%;
  border-collapse: collapse;
  font-size: 13px;
}}
.quality-table th {{ background: #243B64; color: white; }}
.quality-table th, .quality-table td {{
  padding: 11px;
  border-bottom: 1px solid #E5EAF2;
  text-align: left;
}}
.method {{
  background: #EAF0FB;
  border: 1px solid #D5E1F3;
  padding: 17px 20px;
  border-radius: 14px;
  color: #334155;
  line-height: 1.65;
  font-size: 13px;
}}
footer {{
  color: var(--muted);
  font-size: 12px;
  text-align: center;
  padding: 25px;
}}
@media(max-width: 950px) {{
  .kpi-grid {{ grid-template-columns: repeat(2, minmax(0, 1fr)); }}
  .chart-grid {{ grid-template-columns: 1fr; }}
  .chart-card:last-child {{ grid-column: auto; }}
}}
@media(max-width: 600px) {{
  .kpi-grid, .findings-grid {{ grid-template-columns: 1fr; }}
  header {{ padding: 25px 18px; }}
}}
</style>
</head>
<body>
<header>
  <div class="eyebrow">Portfolio Research · Category 1</div>
  <h1>Market &amp; Asset Diagnostics</h1>
  <p>Interactive historical analysis of asset performance, volatility, correlation,
  drawdowns, return distributions and extreme observations. All findings are
  sample-dependent and are not investment advice or forecasts.</p>
</header>
<main>
  <div class="section-title">
    <div>
      <h2>At a glance</h2>
      <p>Key statistics for the aligned sample.</p>
    </div>
  </div>
  <div class="kpi-grid">{kpi_html}</div>

  <div class="section-title">
    <div>
      <h2>Executive Findings</h2>
      <p>Data-driven summaries of the current historical sample.</p>
    </div>
  </div>
  <div class="findings-grid">{finding_html}</div>

  <div class="section-title">
    <div>
      <h2>Interactive Charts</h2>
      <p>Hover for values. Use the Plotly toolbar to zoom, pan, reset, or export.</p>
    </div>
  </div>
  <div class="chart-grid">{''.join(chart_sections)}</div>

  <div class="section-title">
    <div>
      <h2>Asset-Level Metrics</h2>
      <p>Search, filter, sort columns, and export the visible rows.</p>
    </div>
  </div>

  <div class="table-panel">
    <div class="table-controls">
      <input id="tableSearch" type="search"
             placeholder="Search asset or metric values...">
      <select id="assetFilter">
        <option value="">All assets</option>
        {''.join(f'<option value="{html.escape(a)}">{html.escape(a)}</option>' for a in assets)}
      </select>
      <button class="action-btn" onclick="downloadFilteredCSV()">
        Download filtered CSV
      </button>
      <button class="action-btn secondary" onclick="window.print()">
        Print / Save PDF
      </button>
    </div>
    <div class="table-scroll">{table_html}</div>
    <p id="tableCount"></p>
    <p style="font-size:12px;color:#61718A;line-height:1.6">
      Interpret every metric according to its definition. A more negative
      maximum drawdown is worse, while higher volatility is not automatically
      better. The Excel workbook includes daily series and analysis sheets.
    </p>
  </div>

  <div class="section-title">
    <div>
      <h2>Data-Quality Review</h2>
      <p>Review flags before excluding or transforming any observations.</p>
    </div>
  </div>
  <div class="quality-wrap">{quality_html}</div>

  <div class="section-title">
    <div><h2>Methodology &amp; Interpretation</h2></div>
  </div>
  <div class="method">
    <b>Annualization:</b> arithmetic mean return × 252; volatility × √252.
    <b>CAGR:</b> price-based compounded annual growth over the observed interval.
    <b>Sharpe ratio:</b> (annualized arithmetic mean return − 2% assumed annual
    risk-free rate) / annualized volatility.
    <b>Sortino ratio:</b> excess annualized arithmetic mean divided by
    annualized downside deviation.
    <b>Drawdown:</b> price divided by its running maximum minus one.
    <br><br>
    <b>Important:</b> inputs are aligned to common dates with valid observations
    across the included assets. The extreme-return threshold is a descriptive
    screen, not proof of bad data. Correlation regimes are chronological thirds,
    not statistically inferred market regimes. Historical results do not
    forecast future returns.
  </div>
</main>
<footer>
  Generated from processed project data · Interactive HTML report ·
  Category 1: Market and Asset Diagnostics
</footer>
<script>
const table = document.getElementById('metricsTable');
const tbody = table.querySelector('tbody');
const searchBox = document.getElementById('tableSearch');
const assetFilter = document.getElementById('assetFilter');
const countLabel = document.getElementById('tableCount');

function filterRows() {{
  const q = searchBox.value.toLowerCase().trim();
  const asset = assetFilter.value;
  let visible = 0;

  for (const row of tbody.rows) {{
    const matchesSearch = row.innerText.toLowerCase().includes(q);
    const matchesAsset = !asset || row.dataset.asset === asset;
    const show = matchesSearch && matchesAsset;
    row.style.display = show ? '' : 'none';
    if (show) visible++;
  }}

  countLabel.textContent =
    `Showing ${{visible}} of ${{tbody.rows.length}} assets`;
}}

searchBox.addEventListener('input', filterRows);
assetFilter.addEventListener('change', filterRows);

table.querySelectorAll('th').forEach((th, col) => {{
  function sortColumn() {{
    const rows = Array.from(tbody.rows);
    const ascending = th.dataset.direction !== 'asc';

    table.querySelectorAll('th').forEach(h => {{
      delete h.dataset.direction;
    }});

    th.dataset.direction = ascending ? 'asc' : 'desc';

    rows.sort((a, b) => {{
      const av = a.cells[col].dataset.value || a.cells[col].innerText;
      const bv = b.cells[col].dataset.value || b.cells[col].innerText;
      const an = Number(av);
      const bn = Number(bv);
      let comparison;

      if (
        av.trim() !== '' &&
        bv.trim() !== '' &&
        Number.isFinite(an) &&
        Number.isFinite(bn)
      ) {{
        comparison = an - bn;
      }} else {{
        comparison = av.localeCompare(
          bv,
          undefined,
          {{ numeric: true, sensitivity: 'base' }}
        );
      }}

      return ascending ? comparison : -comparison;
    }});

    rows.forEach(row => tbody.appendChild(row));
    filterRows();
  }}

  th.addEventListener('click', sortColumn);

  th.addEventListener('keydown', event => {{
    if (event.key === 'Enter' || event.key === ' ') {{
      event.preventDefault();
      sortColumn();
    }}
  }});
}});

function downloadFilteredCSV() {{
  const visibleRows = Array.from(tbody.rows).filter(
    row => row.style.display !== 'none'
  );

  const headers = Array.from(
    table.querySelectorAll('thead th')
  ).map(th => th.innerText.replace('↕', '').trim());

  const allRows = [headers];

  for (const row of visibleRows) {{
    allRows.push(
      Array.from(row.cells).map(cell => cell.innerText)
    );
  }}

  const csv = allRows.map(row =>
    row.map(value =>
      '"' + String(value).replaceAll('"', '""') + '"'
    ).join(',')
  ).join('\\r\\n');

  const blob = new Blob(
    ['\\uFEFF' + csv],
    {{ type: 'text/csv;charset=utf-8;' }}
  );

  const url = URL.createObjectURL(blob);
  const link = document.createElement('a');
  link.href = url;
  link.download = 'filtered_asset_metrics.csv';
  document.body.appendChild(link);
  link.click();
  link.remove();
  URL.revokeObjectURL(url);
}}

filterRows();
</script>
</body>
</html>
"""

    HTML_OUT.write_text(page, encoding="utf-8")

    # --------------------------------------------------------
    # 16. COMPLETION SUMMARY
    # --------------------------------------------------------

    print("\nDASHBOARD CREATED")
    print(f"Main dashboard : {HTML_OUT}")
    print(f"Metrics CSV    : {CSV_OUT}")
    print(f"Excel workbook : {XLSX_OUT}")
    print(f"Detail pages   : {DETAILS_DIR}")
    print(f"Aligned rows   : {len(returns):,}")
    print(f"Assets         : {len(assets)}")
    print(f"Extreme flags  : {total_extreme:,}")

    print("\nEXECUTIVE FINDINGS")

    for title, body in findings:
        print(f"\n{title}: {body}")

    print("\nExcel sheets created:")
    print(
        "Asset Metrics, Data Quality, Aligned Prices, Daily Returns, "
        "Normalized Performance, Rolling Volatility, Drawdowns, "
        "Full Correlation, Risk Return, Extreme Observations, "
        "Methodology, Early Correlation, Middle Correlation, "
        "Recent Correlation, Correlation Regimes"
    )


if __name__ == "__main__":
    main()